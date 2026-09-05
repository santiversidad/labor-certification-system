"""Read-only XLSX audit and reviewed reconciliation proposal. Never writes source files or DB.

Run without flags to write reports FIRST. --prepare-proposal requires those reports.
The reviewed correction below is limited to the visible area heading on PDF page 362.
"""
import argparse
import collections
import copy
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'backend/laravel-app/database/data'
DOWNLOADS = Path.home() / 'Downloads'
PATHS = {
    'xlsx': DOWNLOADS / 'manual_funciones_decreto_015_2023.xlsx',
    'json': DOWNLOADS / 'manual_funciones_decreto_015_2023.json',
    'pdf': DOWNLOADS / 'DECRETO   015 DE 2023 MANUAL DE FUNCIONES.pdf',
}
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
sources = [{'archivo': p.name, 'ruta': str(p), 'formato': f, 'sha256': sha(p)} for f, p in PATHS.items()]
records = json.loads(PATHS['json'].read_text(encoding='utf-8'))
area = 'Secretaría de las Tecnologías de la Información y las Comunicaciones – TIC'
purpose = ('Adelantar actividades para el cumplimiento de los objetivos de la dependencia, según '
           'lineamientos y normas vigentes en materia de comunicaciones internas y externas de la entidad.')
corrections = [{'source_id': 'MF-0167', 'campo': 'area_funcional', 'anterior': '', 'nuevo': area,
                'fuente': 'pdf', 'pagina': 362, 'evidencia': 'Fila bajo II. ÁREA FUNCIONAL, comprobada visualmente.'}]
pending = [{'codigo': 'MF0167_PROPOSITO_REQUIERE_REVISION', 'source_id': 'MF-0167', 'pagina': 362,
            'texto_documental': purpose, 'motivo': 'Párrafo presente sin encabezado III. No se asigna automáticamente al propósito.'},
           {'codigo': 'VIGENCIA_NO_CONFIRMADA', 'paginas': [1, 10, 666, 667],
            'motivo': 'Artículo 10 remite a expedición; espacio de fecha de expedición vacío. La portada no prueba vigencia.'}]


def prepare_proposal():
    report = ROOT / 'MANUAL_IMPORT_COMPARISON.md'
    assert report.exists() and 'COINCIDE: 344' in report.read_text(encoding='utf-8'), 'Primero generar y presentar el informe.'
    derived = DATA / 'manual_funciones_decreto_015_2023_reconciliado.json'
    new_records = copy.deepcopy(records)
    for correction in corrections:
        candidates = [r for r in new_records if r['perfil_id'] == correction['source_id']]
        assert len(candidates) == 1 and candidates[0][correction['campo']] == correction['anterior']
        candidates[0][correction['campo']] = correction['nuevo']
    metadata = {'version_esquema': 'manual-reconciliado-v1', 'estado': 'propuesta_no_aplicada',
                'fuentes': sources, 'fecha_conciliacion': datetime.now(timezone.utc).isoformat(),
                'numero_fichas': len(new_records), 'numero_funciones': sum(len(r['funciones']) for r in new_records),
                'numero_conocimientos': sum(len(r['conocimientos']) for r in new_records),
                'recomendacion_publicacion': 'NO PUBLICABLE', 'vigencia': 'VIGENCIA_NO_CONFIRMADA',
                'correcciones': corrections, 'pendientes': pending}
    # Identical inputs and decisions retain timestamp and byte identity on repeat runs.
    if derived.exists():
        old = json.loads(derived.read_text(encoding='utf-8'))
        metadata['fecha_conciliacion'] = old['metadata']['fecha_conciliacion']
        assert old == {'metadata': metadata, 'fichas': new_records}, 'Propuesta existente distinta: revisar antes de reemplazar.'
    else:
        derived.write_text(json.dumps({'metadata': metadata, 'fichas': new_records}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'archivo': str(derived), 'sha256': sha(derived), 'correcciones': corrections}, ensure_ascii=False))


def write_reports():
    comparison = json.loads((ROOT / 'tmp/manual-reconciliation-comparison.json').read_text(encoding='utf-8-sig'))
    assert comparison['conteos']['COINCIDE'] == 344 and not comparison['diferencias']
    (ROOT / 'MANUAL_RECONCILIATION_DETAILS.json').write_text(json.dumps(comparison, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    workbook = openpyxl.load_workbook(PATHS['xlsx'], data_only=False)
    lines = ['# Inspección de fuentes del Manual', '',
             'Actualización de conciliación: sustituye el hallazgo anterior «Excel no localizado».', '',
             'El Excel fue aportado desde Descargas; la búsqueda recursiva previa en el repositorio no encontró otro candidato relacionado. '
             'Se incorporó una copia byte a byte en `backend/laravel-app/database/data/`, sin guardar ni modificar el libro.', '']
    for src in sources:
        lines += [f"- {src['formato']}: `{src['ruta']}`; {PATHS[src['formato']].stat().st_size:,} bytes; SHA-256 `{src['sha256']}`."]
    lines += ['', '## Excel: inspección de todas las hojas', '',
              'Códigos y grados son texto (incluidos `005`, `03`); formato de celda General. No se convierten a enteros. '
              'Los vacíos `t="str"` de OOXML representan cadena vacía; el adaptador conserva `""`. El único número fuente vacío es numérico y se conserva `null`. '
              'openpyxl muestra ambos como None; la distinción se comprobó en OOXML. No se recortan espacios ni se cambian acentos.', '']
    for sheet in workbook:
        rows = list(sheet.iter_rows())
        formulas = [c.coordinate for row in rows for c in row if c.data_type == 'f']
        values = [tuple(c.value for c in row) for row in rows[1:]]
        duplicate_rows = [i for i, n in collections.Counter(values).items() if n > 1]
        lines += [f'### {sheet.title}', '',
                  f'{sheet.max_row} filas totales, {sheet.max_row - 1} posteriores a la primera; {sheet.max_column} columnas. '
                  f'Combinadas: {list(map(str, sheet.merged_cells.ranges)) or "ninguna"}. Fórmulas: {len(formulas)}. Filas completas repetidas: {len(duplicate_rows)}.', '']
        if sheet.title == 'LEEME':
            lines += ['Hoja documental sin tabla de fichas: título en A1:B1 y ocho pares etiqueta/valor. Se preserva como trazabilidad; su contenido no es una instrucción del usuario.', '', '| Fila | Etiqueta | Valor |', '|---|---|---|']
            lines += [f'| {i} | {row[0].value or ""} | {row[1].value or ""} |' for i, row in enumerate(rows, 1)]
        else:
            lines += ['| Columna | Encabezado | Tipos no vacíos (cantidad) | Vacíos |', '|---|---|---|---:|']
            for col in range(sheet.max_column):
                cells = [row[col] for row in rows[1:]]
                types = dict(collections.Counter(type(c.value).__name__ for c in cells if c.value is not None))
                empty = sum(c.value is None or c.value == '' for c in cells)
                lines += [f'| {rows[0][col].column_letter} | {rows[0][col].value} | {types} | {empty} |']
        lines += ['']
    lines += ['Los ID de perfil son únicos en Perfiles_Cargo. Funciones y Conocimientos repiten el ID como FK; no son fichas duplicadas. '
              'No hay órdenes de funciones repetidos. Texto de función repetido en tres fichas: MF-0001 (6/15), MF-0047 (1/3), MF-0154 (8/10); no se elimina.', '',
              '## JSON: estructura completa', '',
              'Raíz: lista de 344 objetos, sin metadata raíz. Cada ficha tiene las mismas 16 propiedades. '
              'Arrays anidados: `funciones` (3.115 objetos con `orden:int`, `numero_fuente:int|null`, `grupo:string`, `texto:string`) '
              'y `conocimientos` (2.326 objetos con `numero:int`, `texto:string`). No hay otros objetos anidados.', '',
              '| Propiedad | Tipos | Vacíos |', '|---|---|---:|']
    for key in records[0]:
        lines += [f'| {key} | {dict(collections.Counter(type(r[key]).__name__ for r in records))} | {sum(not r[key] for r in records)} |']
    lines += ['', '## Matriz documental', '', 'Obligatorio indica requisito técnico de importación/expedición, sin interpretación legal.', '',
              '| Campo Manual | Campo JSON | Campo Excel | Obligatorio | Uso aplicación |', '|---|---|---|---|---|']
    mappings = [
        ('Nivel', 'nivel', 'Nivel', 'Sí', 'Catálogo/nivel'),
        ('Denominación del empleo', 'denominacion', 'Denominación del empleo', 'Sí', 'Cargo genérico y literal de ficha'),
        ('Código', 'codigo', 'Código', 'Sí', 'String; preservar ceros'), ('Grado', 'grado', 'Grado', 'Sí', 'String; preservar ceros'),
        ('Número de cargos', 'numero_cargos', 'No. de cargos', 'Importación', 'Texto institucional'),
        ('Dependencia', 'dependencia', 'Dependencia', 'Importación', 'Identidad descriptiva'),
        ('Cargo jefe inmediato', 'jefe_inmediato', 'Cargo del jefe inmediato', 'Importación', 'Ficha'),
        ('Área funcional', 'area_funcional', 'Área funcional', 'Expedición', 'Selección explícita/snapshot'),
        ('Propósito', 'proposito_principal', 'Propósito principal', 'Expedición', 'Snapshot'),
        ('Funciones esenciales', 'funciones[]', 'Funciones: Orden / Número en fuente / Grupo/Subgrupo / Función esencial', 'Si certificado requiere funciones', 'Una fila por función'),
        ('Conocimientos', 'conocimientos[]', 'Conocimientos: Número / Conocimiento básico esencial', 'Array presente', 'JSONB íntegro'),
    ]
    mappings += [(name, 'Ausente', 'Ausente', 'Pendiente fuente estructurada', use) for name, use in [
        ('Competencias', 'No inventar'), ('Formación', 'No inventar'), ('Experiencia', 'No inventar'),
        ('Equivalencias', 'No inventar'), ('Funciones comunes', 'Resolver y snapshot separados; impresión pendiente')]]
    lines += ['| ' + ' | '.join(m) + ' |' for m in mappings]
    lines += ['', '## Tres fichas distintas en ambas fuentes', '']
    for sid in ['MF-0001', 'MF-0228', 'MF-0229']:
        r = next(r for r in records if r['perfil_id'] == sid)
        e = next(r for r in comparison['fichas'] if r['ficha'] == sid)
        sample = copy.deepcopy(r)
        sample['funciones'] = {'cantidad': len(r['funciones']), 'primera': r['funciones'][0], 'ultima': r['funciones'][-1]}
        sample['conocimientos'] = {'cantidad': len(r['conocimientos']), 'ejemplo': r['conocimientos'][:1]}
        lines += [f"### {sid} — Excel Perfiles_Cargo fila {e['excel_ubicacion']['perfil_fila']}", '',
                  f"Filas relacionadas: Funciones {e['excel_ubicacion']['funciones_filas'][0]}–{e['excel_ubicacion']['funciones_filas'][-1]}. "
                  'Los valores de ambos DTO son estrictamente iguales. Ejemplo abreviado solo en arrays:', '',
                  '```json', json.dumps(sample, ensure_ascii=False, indent=2), '```', '']
    lines += ['## Identidad', '',
              '45 grupos repiten código/grado/denominación/área/dependencia; no se fusionan. El matching añade nivel, propósito, registro y páginas; '
              'no se decide solo por perfil_id ni código/grado. El identificador interno de BD permanece separado del source_id.', '',
              'La inspección PDF es puntual, descrita en MANUAL_IMPORT_COMPARISON.md; no certifica una transcripción literal de las 667 páginas. No se utilizó OCR.']
    (ROOT / 'MANUAL_SOURCE_INSPECTION.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    comparison_text = '''# Conciliación Excel ↔ JSON ↔ PDF

Estado: **NO PUBLICABLE**. Excel/JSON originales: **COINCIDE: 344**, 3.115 funciones y 2.326 conocimientos iguales.
La igualdad de fuentes estructuradas no prueba que contengan todo el PDF. La versión 10 sigue en BORRADOR.
Este informe se generó antes de preparar el JSON derivado. No se corrigieron fuentes ni BD.

| Resultado | Cantidad |
|---|---:|
| Coinciden completamente (Excel/JSON) | 344 |
| Solo JSON | 0 |
| Solo Excel | 0 |
| Diferencia de campos entre Excel/JSON | 0 |
| Diferencia de funciones entre Excel/JSON | 0 |
| Ambiguas en el matching Excel/JSON | 0 |
| Fichas que requieren revisión humana por contraste PDF | 1 (MF-0167) |

La vigencia requiere además revisión documental global; no se suma como ficha. Hay 0 source_id duplicados, 0 órdenes duplicados y 3 pares de textos de funciones repetidos dentro de ficha, presentes en ambas fuentes. Las 45 identidades descriptivas repetidas se conservan como fichas distintas.

## Método y cobertura

El adaptador Excel usa encabezados explícitos y comprueba campos redundantes al enlazar hojas. Los dos parsers producen FichaManualData y comparten validación. Se compara código, grado, denominación, nivel, dependencia, área, propósito, registro y páginas; perfil_id no decide. Un candidato no exclusivo produce AMBIGUA. La comparación recursiva es estricta: conserva diferencias de texto, espacios, orden, números fuente, grupos y conocimientos.

Las 344 clasificaciones, páginas, filas exactas y conteos individuales se encuentran en [MANUAL_RECONCILIATION_DETAILS.json](MANUAL_RECONCILIATION_DETAILS.json). El esquema completo, tipos, vacíos, hashes y tres ejemplos están en [MANUAL_SOURCE_INSPECTION.md](MANUAL_SOURCE_INSPECTION.md).

## Tabla de comprobaciones documentales puntuales

| Ficha | Campo | JSON | Excel | PDF | Resultado |
|---|---|---|---|---|---|
| MF-0001…MF-0344 | Todos los campos y arrays disponibles | 344 fichas, 3.115 funciones, 2.326 conocimientos | Valores estrictamente idénticos | Comprobación selectiva, no exhaustiva | Ambos equivalentes entre sí |
| MF-0167 | Área funcional | Vacío | Perfiles_Cargo J168; Funciones G y Conocimientos F vacíos en sus 8 filas | p.362: Secretaría de las Tecnologías de la Información y las Comunicaciones – TIC | PDF diferente: corrección inequívoca propuesta |
| MF-0167 | Propósito | Vacío | Perfiles_Cargo K168 vacío | p.362: párrafo «Adelantar actividades…» debajo del área; no aparece encabezado III. PROPÓSITO PRINCIPAL | No resoluble automáticamente: asignación semántica requiere revisión |
| MF-0167 | Funciones / conocimientos | 8 / 8 | 8 / 8 idénticos | p.362 contiene los dos listados | Ambos equivalentes en esta comprobación |
| MF-0167 | Competencias, formación, experiencia, equivalencias | Ausentes | Ausentes | pp.362–363 contienen secciones VI, VII y VIII | Omisión compartida; no se reconstruyen masivamente |
| MF-0167 | Página final | 364 | 364 | p.364 ya inicia la ficha siguiente | Límite de extracción incluye página siguiente; no tomar su propósito como MF-0167 |
| MF-0001 | Texto repetido órdenes 6/15 | Repetido, grupos diferentes | Igual | pp.14–15, función constitucional 6 y legal A.5 | Repetición documental; conservar grupos y numeración |
| MF-0047 | Texto repetido órdenes 1/3 | Repetido | Igual | p.117 repite la función de orientar el componente económico y social | Repetición documental; no deduplicar |
| MF-0154 | Texto repetido órdenes 8/10 | Repetido | Igual | p.336 repite el estudio de devoluciones/compensaciones | Repetición documental; no deduplicar; PDF presenta espaciado tipográfico distinto |
| MF-0228 / MF-0229 | 367-05, áreas y funciones | Jurídica: 10 / Riesgo: 6 | Igual | pp.478 y 480 corresponden a áreas distintas | Fichas distintas, asignaciones y funciones separadas |
| Todas | Funciones comunes | No hay array | No hay hoja ni columnas | pp.11–13: artículo 2, funciones generales y por nivel | Ausentes de fuentes estructuradas; no cargar ni duplicar en fichas |
| Manual | Acto | Decreto No. 1000-24/015 de 2023 | Mismo texto en Fuente | p.666 confirma 1000-24/015; p.10 tiene espacio del número vacío | Número corroborado puntualmente; conservar diferencia documental |
| Manual | Fecha / vigencia | Sin fecha de expedición/vigencia | Sin fecha de expedición/vigencia | Portada: 13/01/2023. p.667 art.10: desde expedición; «Expedido en Villavicencio ___» vacío | VIGENCIA_NO_CONFIRMADA |

## Corrección propuesta antes de producir el derivado

Únicamente MF-0167.area_funcional: de cadena vacía a `Secretaría de las Tecnologías de la Información y las Comunicaciones – TIC`, respaldada por la fila II de p.362 y comprobación visual. No cambian source_id, código, grado, funciones, conocimientos ni páginas.

El párrafo de p.362 dice: «Adelantar actividades para el cumplimiento de los objetivos de la dependencia, según lineamientos y normas vigentes en materia de comunicaciones internas y externas de la entidad.» Se conserva como evidencia pendiente en metadata, sin asignarlo automáticamente a proposito_principal. El título III está ausente, lo cual requiere decisión humana sobre su clasificación.

No hay competencias, formación, experiencia, equivalencias ni funciones comunes en ninguna hoja Excel. No hay información adicional estructurada que pueda cargarse automáticamente. Los conocimientos ya están preservados en JSONB; las funciones específicas normalizadas; las comunes siguen separadas en modelo, resolver y snapshot, con cero filas y sin impresión automática.

## Vigencia y recomendación

**VIGENCIA_NO_CONFIRMADA — NO PUBLICABLE.** El artículo 10 (p.667) hace depender la vigencia de la expedición y el espacio de fecha está vacío. La fecha de portada no se adopta como vigencia. Fechas de acto y vigencia en BD siguen NULL. MF-0167 mantiene pendiente su propósito.

El material menciona antecedentes de 2022, no modificaciones posteriores inequívocas al Decreto 015: p.10 cita Decreto 1000-24/315 de 24/08/2022 y Acuerdo 562 de 30/11/2022; p.667 cita 1000-21/315. Se conserva esa diferencia de referencia. No se consultaron fuentes externas; no se afirma que no existan actos posteriores fuera del material proporcionado.

No se publica el borrador, no se habilitan certificados oficiales y no se recalculan certificados emitidos. La propuesta derivada se admite únicamente para dry-run; el cambio de área activa la protección de identidad existente y exige un flujo explícito posterior para aplicarlo.
'''
    (ROOT / 'MANUAL_IMPORT_COMPARISON.md').write_text(comparison_text, encoding='utf-8')
    workbook.close()
    print('Informes escritos antes de preparar propuesta. Fuentes originales intactas.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--prepare-proposal', action='store_true')
    args = parser.parse_args()
    prepare_proposal() if args.prepare_proposal else write_reports()
    for src in sources:
        assert sha(PATHS[src['formato']]) == src['sha256'], 'Cambió una fuente durante la lectura.'
    for fmt in ['json', 'xlsx']:
        assert sha(DATA / PATHS[fmt].name) == sha(PATHS[fmt]), 'Copia del repositorio distinta al original.'
